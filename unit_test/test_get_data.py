import unittest
from unittest.mock import patch, Mock
import io
from reportlab.pdfgen import canvas
from docx import Document
from openpyxl import Workbook

from modules.get_data_module import get_data


class TestGetData(unittest.TestCase):
    def generate_pdf_blob(self) -> bytes:
        """Generate a PDF blob with sample text using ReportLab."""
        buffer = io.BytesIO()
        c = canvas.Canvas(buffer)
        c.drawString(10, 10, "Sample PDF text")
        c.save()
        return buffer.getvalue()

    def generate_docx_blob(self) -> bytes:
        """Generate a DOCX blob with sample text using python-docx."""
        buffer = io.BytesIO()
        doc = Document()
        doc.add_paragraph("Sample DOCX text")
        doc.save(buffer)
        return buffer.getvalue()

    def generate_xlsx_blob(self) -> bytes:
        """Generate an XLSX blob with sample text using openpyxl."""
        buffer = io.BytesIO()
        wb = Workbook()
        ws = wb.active
        ws["A1"] = "Sample XLSX text"
        wb.save(buffer)
        return buffer.getvalue()

    @patch("requests.get")
    def test_get_data(self, mock_get: Mock) -> None:
        # Mock the URL response
        mock_response = Mock()
        mock_response.content = b"<html><body>Sample URL text</body></html>"
        mock_response.headers = {"Content-Type": "text/html"}
        mock_response.raise_for_status.return_value = None
        mock_get.return_value = mock_response

        # Generate test file blobs
        pdf_blob = self.generate_pdf_blob()
        docx_blob = self.generate_docx_blob()
        xlsx_blob = self.generate_xlsx_blob()

        # Define test inputs
        files = [pdf_blob, docx_blob, xlsx_blob]
        urls = ["http://example.com"]

        # Call the function
        result = get_data(files, urls)

        # Assert extracted text contains expected values
        self.assertIn("Sample PDF text", result)
        self.assertIn("Sample DOCX text", result)
        self.assertIn("Sample XLSX text", result)
        self.assertIn("Sample URL text", result)

if __name__ == "__main__":
    unittest.main()